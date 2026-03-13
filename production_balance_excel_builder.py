from openpyxl.styles import Font, PatternFill, Alignment, Border, Side


HEADER_FILL = PatternFill(fill_type="solid", fgColor="D9EAF7")
HEADER_FONT = Font(bold=True)

THIN_BORDER = Border(
    left=Side(style="thin"),
    right=Side(style="thin"),
    top=Side(style="thin"),
    bottom=Side(style="thin"),
)


def style_header_row(ws, row_num: int):
    for cell in ws[row_num]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = THIN_BORDER


def apply_borders(ws):
    for row in ws.iter_rows(
        min_row=1,
        max_row=ws.max_row,
        min_col=1,
        max_col=ws.max_column,
    ):
        for cell in row:
            if cell.value is not None and cell.value != "":
                cell.border = THIN_BORDER


def autosize_columns(ws, max_width: int = 28):
    for column_cells in ws.columns:
        max_length = 0
        col_letter = column_cells[0].column_letter

        for cell in column_cells:
            if cell.value is not None:
                cell_len = len(str(cell.value))
                if cell_len > max_length:
                    max_length = cell_len

        adjusted_width = min(max_length + 2, max_width)
        ws.column_dimensions[col_letter].width = adjusted_width


def build_production_balance_sheet(wb, production_balance: dict):
    ws = wb.create_sheet("Production Balance")

    headers = [
        "Ресурс",
        "Производство / час",
        "Производство / сутки",
        "Потребление / час",
        "Потребление / сутки",
        "Баланс / сутки",
    ]

    ws.append(headers)
    style_header_row(ws, 1)

    for row in production_balance["balance_rows"]:
        if row["group"] == "separator":
            ws.append(["", "", "", "", "", ""])
            continue

        ws.append([
            row["resource_name"],
            round(row["production_hour"], 2),
            round(row["production_day"], 2),
            round(row["consumption_hour"], 2),
            round(row["consumption_day"], 2),
            round(row["balance_day"], 2),
        ])

    # Выравнивание
    for row_num in range(2, ws.max_row + 1):
        for col_num in range(1, 7):
            cell = ws.cell(row=row_num, column=col_num)
            cell.alignment = Alignment(horizontal="center", vertical="center")

    apply_borders(ws)
    autosize_columns(ws, max_width=24)
