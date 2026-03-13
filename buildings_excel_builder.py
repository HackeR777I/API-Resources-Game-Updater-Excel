from openpyxl.styles import Font, PatternFill, Alignment, Border, Side


HEADER_FILL = PatternFill(fill_type="solid", fgColor="D9EAF7")
HEADER_FONT = Font(bold=True)

THIN_BORDER = Border(
    left=Side(style="thin"),
    right=Side(style="thin"),
    top=Side(style="thin"),
    bottom=Side(style="thin"),
)


def style_header_row(ws, row_num: int):
    for cell in ws[row_num]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = THIN_BORDER


def apply_borders(ws):
    for row in ws.iter_rows(
        min_row=1,
        max_row=ws.max_row,
        min_col=1,
        max_col=ws.max_column,
    ):
        for cell in row:
            if cell.value is not None:
                cell.border = THIN_BORDER


def autosize_columns(ws, max_width: int = 35):
    for column_cells in ws.columns:
        max_length = 0
        col_letter = column_cells[0].column_letter

        for cell in column_cells:
            if cell.value is not None:
                cell_len = len(str(cell.value))
                if cell_len > max_length:
                    max_length = cell_len

        adjusted_width = min(max_length + 2, max_width)
        ws.column_dimensions[col_letter].width = adjusted_width


# ---------------------------
# FACTORIES
# ---------------------------

def build_factories_sheet(wb, factory_stats: dict):
    ws = wb.create_sheet("Factories")

    headers = [
        "Завод",
        "Уровень",
        "Продукт",
        "Производство / час",
        "Производство / сутки",
        "Потребление ресурсов",
    ]

    ws.append(headers)
    style_header_row(ws, 1)

    for row in factory_stats["factory_rows"]:

        inputs_text = "\n".join(
        f'{i["item_name"]} ({round(i["amount_per_hour"], 2)}/ч)'
        for i in row["inputs"]
        )

        ws.append([
            row["factory_name"],
            row["level"],
            row["output_item_name"],
            round(row["output_per_hour"], 2),
            round(row["output_per_day"], 2),
            inputs_text
        ])

    for row_num in range(2, ws.max_row + 1):
        ws.cell(row=row_num, column=6).alignment = Alignment(wrap_text=True, vertical="top")
        ws.row_dimensions[row_num].height = 35
    apply_borders(ws)
    autosize_columns(ws)


# ---------------------------
# SPECIAL BUILDINGS
# ---------------------------

def build_special_buildings_sheet(wb, special_building_stats: dict):
    ws = wb.create_sheet("Special Buildings")

    headers = [
        "Здание",
        "Уровень",
    ]

    ws.append(headers)
    style_header_row(ws, 1)

    for row in special_building_stats["special_building_rows"]:
        ws.append([
            row["building_name"],
            row["level"],
        ])

    apply_borders(ws)
    autosize_columns(ws)


# ---------------------------
# WAREHOUSES
# ---------------------------

def build_warehouses_sheet(wb, warehouse_stats: dict):
    ws = wb.create_sheet("Warehouses")

    headers = [
        "Ресурс",
        "Уровень склада",
        "Хранится",
        "Вместимость",
    ]

    ws.append(headers)
    style_header_row(ws, 1)

    for row in warehouse_stats["warehouse_rows"]:
        ws.append([
            row["resource_name"],
            row["level"],
            round(row["stored_amount"], 2),
            round(row["capacity"], 2),
        ])

    apply_borders(ws)
    autosize_columns(ws)


# ---------------------------
# MASTER FUNCTION
# ---------------------------

def build_buildings_workbook_part(
    wb,
    factory_stats: dict,
    special_building_stats: dict,
    warehouse_stats: dict
):
    build_factories_sheet(wb, factory_stats)
    build_special_buildings_sheet(wb, special_building_stats)
    build_warehouses_sheet(wb, warehouse_stats)
