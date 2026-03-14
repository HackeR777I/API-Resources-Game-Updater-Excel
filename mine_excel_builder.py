from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side


HEADER_FILL = PatternFill(fill_type="solid", fgColor="D9EAF7")
HEADER_FONT = Font(bold=True)

THIN_BORDER = Border(
    left=Side(style="thin"),
    right=Side(style="thin"),
    top=Side(style="thin"),
    bottom=Side(style="thin"),
)


def style_header_row(ws, row_num: int):
    for cell in ws[row_num]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = THIN_BORDER


def apply_borders(ws):
    for row in ws.iter_rows(
        min_row=1,
        max_row=ws.max_row,
        min_col=1,
        max_col=ws.max_column,
    ):
        for cell in row:
            if cell.value is not None:
                cell.border = THIN_BORDER


def autosize_columns(ws, max_width: int = 30):
    for column_cells in ws.columns:
        max_length = 0
        col_letter = column_cells[0].column_letter

        for cell in column_cells:
            if cell.value is not None:
                cell_len = len(str(cell.value))
                if cell_len > max_length:
                    max_length = cell_len

        adjusted_width = min(max_length + 2, max_width)
        ws.column_dimensions[col_letter].width = adjusted_width


def build_mine_summary_sheet(wb, mine_stats: dict):
    ws = wb.create_sheet("Mine Stats")

    headers = [
        "Ресурс",
        "Кол-во шахт",
        "Среднее качество",
        "Мин. качество",
        "Макс. качество",
        "Среднее состояние",
        "Суммарный rawrate / час",
        "Суммарный fullrate / час",
        "Эффективная добыча / час",
        "Эффективная добыча / сутки",
    ]

    ws.append(headers)
    style_header_row(ws, 1)

    for row in mine_stats["summary_rows"]:
        ws.append([
            row["resource_name"],
            row["mine_count"],
            round(row["avg_quality"], 4),
            round(row["min_quality"], 4),
            round(row["max_quality"], 4),
            round(row["avg_condition"], 6),
            round(row["total_rawrate_hour"], 2),
            round(row["total_fullrate_hour"], 2),
            round(row["total_effective_output_hour"], 2),
            round(row["total_effective_output_day"], 2),
        ])

    apply_percent_columns(ws, [3, 4, 5, 6,], start_row=2)
    apply_borders(ws)
    autosize_columns(ws, max_width=28)
    
def apply_percent_columns(ws, columns: list[int], start_row: int = 2):
    for col in columns:
        for row in range(start_row, ws.max_row + 1):
            cell = ws.cell(row=row, column=col)
            if isinstance(cell.value, (int, float)):
                cell.number_format = "0.00%"


def safe_sheet_title(title: str, max_len: int = 31) -> str:
    bad_chars = ['\\', '/', '*', '?', ':', '[', ']']
    for ch in bad_chars:
        title = title.replace(ch, "_")
    return title[:max_len]


def build_mine_resource_sheets(wb, mine_stats: dict):
    resource_sheets = mine_stats["resource_sheets"]

    for resource_name, rows in resource_sheets.items():
        ws = wb.create_sheet(safe_sheet_title(resource_name))

        headers = [
            "№",
            "HQ boost",
            "Tech factor",
            "Raw rate / час",
            "Full rate / час",
            "Condition",
            "Quality",
            "Дата обслуживания",
            "Дата постройки",
            "Эффективная добыча / час",
            "Эффективная добыча / сутки",
        ]

        ws.append(headers)
        style_header_row(ws, 1)

        for row in rows:
            ws.append([
                row["number"],
                round(row["hq_boost"], 4),
                round(row["tech_factor"], 4),
                round(row["rawrate"], 2),
                round(row["fullrate"], 2),
                round(row["condition"], 6),
                round(row["quality"], 4),
                row["last_maintenance"],
                row["build_date"],
                round(row["effective_output_hour"], 2),
                round(row["effective_output_day"], 2),
        ])

        apply_percent_columns(ws, [6, 7], start_row=2)

        apply_borders(ws)
        autosize_columns(ws, max_width=26)


def build_mine_workbook_part(wb, mine_stats: dict):
    build_mine_summary_sheet(wb, mine_stats)
    build_mine_resource_sheets(wb, mine_stats)