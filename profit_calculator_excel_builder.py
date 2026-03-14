from openpyxl.styles import Font, PatternFill, Alignment, Border, Side


HEADER_FILL = PatternFill(fill_type="solid", fgColor="D9EAF7")
HEADER_FONT = Font(bold=True)

THIN_BORDER = Border(
    left=Side(style="thin"),
    right=Side(style="thin"),
    top=Side(style="thin"),
    bottom=Side(style="thin"),
)


def style_header_row(ws, row_num: int):
    for cell in ws[row_num]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
            wrap_text=True
        )
        cell.border = THIN_BORDER


def apply_borders(ws):
    for row in ws.iter_rows(
        min_row=1,
        max_row=ws.max_row,
        min_col=1,
        max_col=ws.max_column,
    ):
        for cell in row:
            if cell.value is not None and cell.value != "":
                cell.border = THIN_BORDER


def set_column_widths(ws):
    widths = {
        "A": 42,
        "B": 24,
    }

    for col, width in widths.items():
        ws.column_dimensions[col].width = width


def apply_number_format(ws, start_row: int, end_row: int, value_col: int = 2):
    for row_num in range(start_row, end_row + 1):
        cell = ws.cell(row=row_num, column=value_col)

        if isinstance(cell.value, (int, float)):
            cell.number_format = '#,##0.00'


def build_profit_summary_sheet(wb, daily_profit_report: dict):
    ws = wb.create_sheet("Profit")

    ws.append(["Показатель", "Значение"])
    style_header_row(ws, 1)

    rows = [
        ["Логистика", daily_profit_report["logistics_percent"]],
        ["Доход по ресурсам / сутки", daily_profit_report["mine_income_day"]],
        ["Расход по ресурсам / сутки", daily_profit_report["mine_expense_day"]],
        ["Чистая прибыль по ресурсам (оценка)", daily_profit_report["mine_net_profit_day"]],
        ["Доход по товарам / сутки", daily_profit_report["factory_income_day"]],
        ["Расход по товарам / сутки", daily_profit_report["factory_expense_day"]],
        ["Чистая прибыль по товарам (оценка)", daily_profit_report["factory_net_profit_day"]],
        ["Стоимость запуска заводов / сутки", daily_profit_report["total_credits_cost_per_day"]],
        ["Итоговая чистая прибыль (оценка)", daily_profit_report["total_net_profit_day"]],
    ]

    for row in rows:
        ws.append(row)

    # Логистика как проценты
    ws["B2"].number_format = "0.00%"

    # Все остальные значения — обычные большие числа
    apply_number_format(ws, start_row=3, end_row=ws.max_row, value_col=2)

    # Выравнивание
    for row_num in range(2, ws.max_row + 1):
        ws.cell(row=row_num, column=1).alignment = Alignment(
            horizontal="left",
            vertical="center",
            wrap_text=True
        )
        ws.cell(row=row_num, column=2).alignment = Alignment(
            horizontal="right",
            vertical="center"
        )

    apply_borders(ws)
    set_column_widths(ws)